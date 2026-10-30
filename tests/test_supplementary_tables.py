"""CSV/TSV/XLSX supplementary tables → records (no network).

uv run --extra supplementary pytest tests
"""

import io
import itertools
import sys
from contextlib import contextmanager
from unittest import mock

import openpyxl

from librarian import supplementary
from librarian.supplementary import extract_supplementary_records


def _xlsx(sheets):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_csv_repeats_header_in_every_record():
    rows = ["gene,log2FC,padj"] + [f"GENE{i},{i}.5,0.01" for i in range(60)]
    records = extract_supplementary_records(
        {"s1.csv": "\n".join(rows).encode()}, {}, max_records=40
    )
    assert len(records) > 1
    for record in records:
        assert record["text"].startswith("gene | log2FC | padj ; ")
    assert "GENE59 | 59.5 | 0.01" in records[-1]["text"]


def test_tsv_uses_tab_delimiter():
    records = extract_supplementary_records(
        {"s1.tsv": b"strain\tgenotype\nMG1655\twild type\n"}, {}, max_records=40
    )
    assert records[0]["text"] == "strain | genotype ; MG1655 | wild type"


def test_xlsx_titles_records_with_sheet_and_caps_them():
    data = _xlsx(
        {
            "Hits": [["gene", "score"]] + [[f"G{i}", i * 1.0] for i in range(500)],
            "Primers": [
                ["name", "sequence"],
                ["SOX2-F", "GGCAGCTACAGCATGATGCAGGA"],
            ],
        }
    )
    records = extract_supplementary_records({"table_s1.xlsx": data}, {}, max_records=5)
    assert len(records) == 5
    assert records[0]["section_title"].endswith(', sheet "Hits"')
    assert "G0 | 0" in records[0]["text"]  # 0.0 shown as 0
    # The cap is shared across sheets, so the short second sheet is read too.
    assert 'sheet "Primers"' in " ".join(r["section_title"] for r in records)


def test_empty_cells_keep_their_column():
    records = extract_supplementary_records(
        {"s1.csv": b"gene,log2FC,padj\nA,,0.01\n,,\nB,2,\n"}, {}, max_records=40
    )
    assert records[0]["text"] == "gene | log2FC | padj ; A |  | 0.01 ; B | 2"


def test_title_row_is_a_caption_and_the_real_header_starts_every_block():
    rows = ["Table S1. DE genes", "gene,log2FC,padj"]
    rows += [f"GENE{i},{i}.5,0.01" for i in range(60)]
    records = extract_supplementary_records(
        {"s1.csv": "\n".join(rows).encode()}, {}, max_records=40
    )
    assert len(records) > 1
    for record in records:
        assert record["text"].startswith("Table S1. DE genes ; gene | log2FC | padj ; ")


def test_rows_are_read_lazily_and_stop_at_the_cap():
    def endless_rows():
        yield ["gene", "score"]
        for i in itertools.count():
            if i > 1000:
                raise AssertionError("read past the record cap")
            yield [f"G{i}", i]

    @contextmanager
    def reader(data):
        yield [("Hits", endless_rows())]

    with mock.patch.dict(supplementary._SHEET_READERS, {".xlsx": reader}):
        records = extract_supplementary_records({"big.xlsx": b""}, {}, max_records=3)
    assert len(records) == 3


def _stage2_pool(files, subquery="query", everything=False, **config_overrides):
    """Every supplementary chunk the agent hands to BM25 for one sub-query.

    Runs the agent's real Stage-2 path (download mocked, then extraction,
    table re-cutting and chunking) and captures the pool before ranking.
    """
    from dataclasses import replace

    from librarian import LibrarianAgent, load_runtime_config
    from librarian.literature_search import EuropePmcSource, Fulltext

    xml = (
        "<article><body><p>Unrelated body text.</p></body><back>"
        "<supplementary-material/></back></article>"
    )
    paper = {
        "title": "T",
        "isOpenAccess": True,
        "pmcid": "PMC1",
        "fullTextIds": ["PMC1"],
        "pmid": "1",
        "year": "2024",
    }

    class Source(EuropePmcSource):
        def search(self, query, page_size):
            return [dict(paper)]

        def fetch_fulltext_many(self, pmcids):
            return {"PMC1": Fulltext("PMC1", xml=xml, status=200)}

    listing = (
        '<ListBucketResult xmlns="http://s3.amazonaws.com/doc/2006-03-01/">'
        + "".join(
            f"<Contents><Key>PMC1.1/{name}</Key><Size>{len(data)}</Size></Contents>"
            for name, data in files.items()
        )
        + "</ListBucketResult>"
    ).encode()

    def get(url, **kwargs):
        response = mock.MagicMock()
        response.content = (
            listing if url.endswith("/") else files[url.rsplit("/", 1)[1]]
        )
        return response

    pool = []

    def capture(subquery, records, k, ranking_terms=()):
        pool.extend(records)
        return []

    config = replace(
        load_runtime_config(), supplementary_enrichment=True, **config_overrides
    )
    agent = LibrarianAgent(config, llm_client=object(), literature_source=Source())
    with (
        mock.patch("requests.get", side_effect=get),
        mock.patch(
            "librarian.agent._rank_paragraphs_for_subquery", side_effect=capture
        ),
    ):
        agent._paragraphs_for_subquery(subquery)
    if everything:
        return config.max_paragraph_words, pool
    return config.max_paragraph_words, [
        r for r in pool if r["section_type"] == "supplementary"
    ]


def _assert_every_chunk(chunks, max_words, startswith):
    assert len(chunks) > 1
    for chunk in chunks:
        assert len(chunk["text"].split()) <= max_words, chunk["text"][:80]
        assert chunk["text"].startswith(startswith), chunk["text"][:80]


# The header invariant on every chunk before BM25, not just BM25's picks.
def test_a_long_row_keeps_header_and_row_identity_in_every_chunk():
    description = " ".join(f"word{i}" for i in range(600))
    max_words, chunks = _stage2_pool(
        {"s1.csv": f"gene,description\nSOX2,{description}\n".encode()}
    )
    _assert_every_chunk(chunks, max_words, "gene | description ; SOX2 | ")
    assert "word599" in " ".join(c["text"] for c in chunks)


def test_a_long_cell_continues_in_its_own_column():
    description = " ".join(f"word{i}" for i in range(600))
    max_words, chunks = _stage2_pool(
        {"s.csv": f"gene,score,description\nSOX2,3,{description}\n".encode()}
    )
    assert len(chunks) > 1
    continued = []
    for chunk in chunks:
        text = chunk["text"]
        assert len(text.split()) <= max_words
        header, row = text.split(" ; ", 1)
        assert header == "gene | score | description"
        cells = row.split(" | ")
        assert len(cells) == 3, text[:180]
        assert cells[0] == "SOX2"
        assert cells[1] in ("3", "")
        assert cells[2].startswith("word")
        continued += cells[2].split()
    assert continued == description.split()  # all of it, in order


def test_a_long_middle_cell_keeps_the_next_column_in_place():
    description = " ".join(f"word{i}" for i in range(600))
    max_words, chunks = _stage2_pool(
        {"s.csv": f"gene,description,padj\nSOX2,{description},0.01\n".encode()}
    )
    padj = []
    for chunk in chunks:
        assert len(chunk["text"].split()) <= max_words
        header, row = chunk["text"].split(" ; ", 1)
        assert header == "gene | description | padj"
        cells = row.split(" | ")
        assert len(cells) == 3, chunk["text"][:180]
        assert cells[0] == "SOX2"
        assert "0.01" not in cells[1]
        padj.append(cells[2].strip())
    assert [p for p in padj if p] == ["0.01"]


def test_rows_packed_together_leave_room_for_every_separator():
    # Three short rows and one 400-word row land in one block; re-cutting it
    # packs whole rows and row parts, each " ; " costing a word.
    small = [f"G{i},{' '.join(['note'] * 20)}" for i in range(3)]
    big = "BIG," + " ".join(f"w{i}" for i in range(400))
    max_words, chunks = _stage2_pool(
        {"s1.csv": ("gene,description\n" + "\n".join(small + [big])).encode()}
    )
    _assert_every_chunk(chunks, max_words, "gene | description ; ")
    assert any(c["text"].count(" ; ") >= 3 for c in chunks)
    assert all(
        c["text"].split(" ; ")[1].split(" | ")[0] in {"G0", "BIG"} for c in chunks
    )


def test_a_wide_table_is_split_into_column_groups_each_with_its_header():
    header = "gene," + ",".join(f"sample{i}" for i in range(130))
    rows = [f"G{r}," + ",".join(str(r * i) for i in range(130)) for r in range(3)]
    max_words, chunks = _stage2_pool(
        {"s1.csv": (header + "\n" + "\n".join(rows)).encode()}
    )
    _assert_every_chunk(chunks, max_words, "gene | sample")
    columns = set()
    for chunk in chunks:
        columns.update(chunk["text"].split(" ; ")[0].split(" | "))
    assert all(f"sample{i}" in columns for i in range(130))


def test_each_sheet_reads_its_own_rows():
    sheets = {
        "Hits": [["gene", "score"], ["SOX2", 3.0]],
        "Primers": [["primer", "sequence"], ["P1", "ACGT"]],
    }

    class Sheet:
        def __init__(self, name, rows):
            self.name, self.rows, self.nrows = name, rows, len(rows)

        def row_values(self, i):
            return self.rows[i]

    fake_xlrd = mock.MagicMock()
    fake_xlrd.open_workbook.return_value.sheets.return_value = [
        Sheet(name, rows) for name, rows in sheets.items()
    ]
    with mock.patch.dict(sys.modules, {"xlrd": fake_xlrd}):
        records = extract_supplementary_records({"t.xls": b""}, {}, max_records=40)
    by_sheet = {
        r["section_title"].split('sheet "')[1].rstrip('"'): r["text"] for r in records
    }
    assert by_sheet == {
        "Hits": "gene | score ; SOX2 | 3",
        "Primers": "primer | sequence ; P1 | ACGT",
    }


def test_docx_tables_keep_empty_cells_and_repeat_the_header():
    import docx

    document = docx.Document()
    document.add_paragraph("Supplementary methods.")
    rows = [["gene", "log2FC", "padj"], ["A", "", "0.01"]]
    rows += [[f"GENE{i}", f"{i}.5", "0.02"] for i in range(60)]
    table = document.add_table(rows=len(rows), cols=3)
    for r, values in enumerate(rows):
        for c, value in enumerate(values):
            table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    records = extract_supplementary_records({"s1.docx": buffer.getvalue()}, {}, 40)
    tables = [r for r in records if r["section_title"].endswith(", table 1")]
    assert len(tables) > 1
    for record in tables:
        assert record["text"].startswith("gene | log2FC | padj ; ")
    assert "A |  | 0.01" in tables[0]["text"]


def test_xlsx_unpacking_past_the_budget_is_not_opened():
    data = _xlsx({"Hits": [["gene", "note"]] + [["G", "x" * 5000]] * 200})
    with (
        mock.patch.object(supplementary, "_MAX_PACKAGE_BYTES", 100_000),
        mock.patch("openpyxl.load_workbook") as load_workbook,
    ):
        records = extract_supplementary_records({"t.xlsx": data}, {}, 40)
    assert records == []
    load_workbook.assert_not_called()
